import json, math, statistics
from pathlib import Path
from datetime import datetime, date
from functools import lru_cache
from urllib.parse import quote
from openpyxl import load_workbook
BASE = Path(__file__).resolve().parent / "runtime"
FILES = {"reviewers_file":"Pareceristas","editorial_dates_file":"Artigos e datas","editorial_board_file":"Conselho editorial","authors_file":"Autores"}
def numeric(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) and value >= 0 else None

def ratio(numerator, denominator):
    if numeric(numerator) is None or numeric(denominator) is None or denominator <= 0 or numerator > denominator:
        return None
    return round(100 * numerator / denominator, 1)

def safe_file(raw):
    if not raw:
        return None
    try:
        candidate = (BASE / raw).resolve(strict=True)
        candidate.relative_to((BASE / "runtime/uploads").resolve())
        return candidate if candidate.is_file() else None
    except (OSError, ValueError):
        return None

def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    return None

def date_inventory(path, modified, size):
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = list(next(rows, []))
            if "titulo_artigo" in headers and "data_recebimento" in headers:
                records.extend(dict(zip(headers, row)) for row in rows if any(v is not None for v in row))
                break
    finally:
        workbook.close()
    return records

def inventory(response, annual):
    path = safe_file(response.get("editorial_dates_file"))
    result = {"available": False, "rows": 0, "unique_articles": 0, "duplicates": 0, "years": [], "annual": []}
    if not path:
        return result
    try:
        stat = path.stat()
        rows = date_inventory(str(path), stat.st_mtime_ns, stat.st_size)
    except Exception:
        return {**result, "error": "Não foi possível ler a planilha de datas."}
    seen = set(); unique = []
    for row in rows:
        doi = str(row.get("doi") or "").lower().strip().removeprefix("https://doi.org/")
        key = doi or (str(row.get("ano")), " ".join(str(row.get("titulo_artigo") or "").lower().split()))
        if key in seen:
            continue
        seen.add(key); unique.append(row)
    def year_of(row):
        try:
            return int(row.get("ano"))
        except (ValueError, TypeError):
            return None
    for item in annual:
        sample = [r for r in unique if year_of(r) == item["year"]]
        deltas = []; invalid = 0
        for row in sample:
            received, accepted = parse_date(row.get("data_recebimento")), parse_date(row.get("data_aprovacao"))
            if received and accepted:
                if accepted >= received:
                    deltas.append((accepted - received).days)
                else:
                    invalid += 1
        ordered = sorted(deltas)
        result["annual"].append({"year": item["year"], "articles": len(sample), "valid_pairs": len(deltas), "invalid_pairs": invalid, "median_days": statistics.median(deltas) if deltas else None, "p90_days": ordered[max(0, math.ceil(len(ordered)*.9)-1)] if deltas else None, "doi_filled": sum(bool(r.get("doi")) for r in sample)})
    result.update(available=True, rows=len(rows), unique_articles=len(unique), duplicates=len(rows)-len(unique), years=sorted({year_of(r) for r in unique if year_of(r) is not None}))
    return result

def submission(row):
    response = json.loads(row["response_json"])
    processed = json.loads(row["processed_json"])
    year = int(response.get("submission_year") or 2026)
    annual = []
    for offset in (2, 1, 0):
        suffix = f"y{offset}"
        published = numeric(response.get(f"published_articles_{suffix}"))
        international = numeric(response.get(f"international_articles_{suffix}"))
        foreign = numeric(response.get(f"foreign_language_articles_{suffix}"))
        annual.append({"year": year-offset, "partial": year-offset >= date.today().year, "published": published, "international": international, "international_pct": ratio(international, published), "foreign": foreign, "foreign_pct": ratio(foreign, published), "received": numeric(response.get(f"received_articles_{suffix}")), "approved": numeric(response.get(f"approved_articles_{suffix}")), "rejected": numeric(response.get(f"rejected_articles_{suffix}"))})
    annexes = []
    for key, label in FILES.items():
        path = safe_file(response.get(key))
        annexes.append({"field": key, "label": label, "exists": bool(path), "name": path.name if path else None, "url": f"/api/annex/{quote(row['response_id'])}/{key}" if path else None})
    alerts = [{"type": "origem", "message": x["message"]} for x in processed.get("warnings", [])]
    for file in annexes:
        if not file["exists"]:
            alerts.append({"type": "anexo", "message": f"Anexo ausente: {file['label']}."})
    missing = sum(item[key] is None for item in annual for key in ("received", "approved", "rejected"))
    if missing:
        alerts.append({"type": "dados", "message": f"{missing} estatísticas de recebidos/aprovados/rejeitados não informadas; taxas indisponíveis."})
    if any(item["partial"] for item in annual):
        alerts.append({"type": "período", "message": f"{year} está em curso. Comparações com anos completos e a janela dos anexos exigem revisão."})
    if response.get("review_time_min_months") == response.get("review_time_max_months") == 0:
        alerts.append({"type": "qualidade", "message": "Prazo declarado de 0 meses: confirmar se o zero representa ausência de informação."})
    if all(item[key] == 0 for item in annual for key in ("received", "approved", "rejected")):
        alerts.append({"type": "qualidade", "message": "Nove zeros nas estatísticas editoriais: confirmar a origem; não interpretar como ausência de submissões."})
    inv = inventory(response, annual)
    if inv.get("error"):
        alerts.append({"type": "anexo", "message": inv["error"]})
    for item in inv["annual"]:
        declared = next(a["published"] for a in annual if a["year"] == item["year"])
        if declared is not None and declared != item["articles"]:
            alerts.append({"type": "conciliação", "message": f"{item['year']}: {declared} artigos declarados e {item['articles']} identificados no anexo. Conferir cobertura e tipos documentais."})
        if item["invalid_pairs"]:
            alerts.append({"type": "datas", "message": f"{item['year']}: {item['invalid_pairs']} datas de aprovação anteriores ao recebimento."})
    criteria = []
    for original in processed.get("criteria", []):
        item = dict(original)
        item["manager_status"] = "Cálculo preliminar" if item.get("status_basis") == "Regra determinística sobre dados declarados" else "Aguardando conferência"
        criteria.append(item)
    return {"id": row["response_id"], "title": row["journal_title"], "issn": response.get("issn_online"), "created_at": row["created_at"], "status": row["status"], "annual": annual, "annexes": annexes, "alerts": alerts, "criteria": criteria, "response": response, "inventory": inv, "missing_stats": missing, "review_declared": [response.get("review_time_min_months"), response.get("review_time_max_months")]}